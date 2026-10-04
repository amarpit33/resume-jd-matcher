"""Generate the fictional sample resume as PDF and DOCX (needs requirements-dev: fpdf2).

    python scripts/make_sample_resumes.py
"""
from __future__ import annotations

from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "data" / "sample_resumes"

HEADER = ["Alex Rivera", "Backend Engineer | alex.rivera@example.com | Pune, India"]
SECTIONS: list[tuple[str, list[str]]] = [
    ("Summary", ["Backend engineer with 6 years building Python services and data pipelines."]),
    ("Skills", ["Languages: Python, SQL, JavaScript",
                "Frameworks: Django, Flask, Celery, pandas",
                "Tools: PostgreSQL, Redis, Docker, Git, AWS (EC2, S3), Jenkins"]),
    ("Experience", [
        "Software Engineer, Finlytics Pvt Ltd | Jun 2022 - Present",
        "- Built and maintained Django REST endpoints for the loan-origination product used by 40 partner banks",
        "- Moved nightly reconciliation jobs from cron scripts to Celery workers, cutting failures by 60%",
        "- Worked on database query tuning in PostgreSQL for slow reporting pages",
        "- Helped the team set up Docker-based local development",
        "- Reviewed pull requests and mentored two interns",
        "Associate Engineer, DataNest Solutions | Jul 2020 - May 2022",
        "- Wrote Flask services to ingest CSV uploads from clients into PostgreSQL",
        "- Created pandas scripts for weekly client reports",
        "- Fixed production bugs and wrote unit tests with pytest",
    ]),
    ("Education", ["B.E. Computer Engineering, Savitribai Phule Pune University, 2020"]),
]


def as_lines() -> list[str]:
    lines = list(HEADER)
    for title, body in SECTIONS:
        lines += ["", title.upper(), *body]
    return lines


def write_pdf(path: Path) -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    for line in as_lines():
        pdf.set_font("Helvetica", "B" if line.isupper() or line == HEADER[0] else "", 10)
        pdf.multi_cell(0, 5, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


def write_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading(HEADER[0], 0)
    d.add_paragraph(HEADER[1])
    for title, body in SECTIONS:
        d.add_heading(title, 1)
        for line in body:
            if line.startswith("- "):
                d.add_paragraph(line[2:], style="List Bullet")
            else:
                d.add_paragraph(line)
    d.save(str(path))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    write_pdf(OUT / "sample_resume.pdf")
    write_docx(OUT / "sample_resume.docx")
    print(f"wrote {OUT}/sample_resume.pdf and .docx")
