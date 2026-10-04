"""Resume ingestion: file bytes -> raw text -> structured sections (LLM)."""
from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass

from api.schemas import ResumeStructured
from api.services.model_client import StructuredLLM
from api.services.storage import Storage

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")


class ResumeParseError(ValueError):
    """The resume file could not be read or contains no usable text."""


@dataclass(frozen=True)
class Bullet:
    """A resume bullet with a stable id the LLM can cite instead of re-quoting it."""

    id: str
    role: str
    text: str


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX, TXT or MD resume.

    Raises:
        ResumeParseError: unsupported type, unreadable file, or no extractable text.
    """
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            text = _pdf_text(data)
        elif name.endswith(".docx"):
            text = _docx_text(data)
        elif name.endswith((".txt", ".md")):
            text = data.decode("utf-8", errors="replace")
        else:
            raise ResumeParseError(f"unsupported file type; use one of {', '.join(SUPPORTED_EXTENSIONS)}")
    except ResumeParseError:
        raise
    except Exception as exc:  # pdfminer/python-docx raise a zoo of exception types
        raise ResumeParseError(f"could not read {filename}: {exc}") from exc
    text = clean_text(text)
    if len(text) < 100:
        raise ResumeParseError(
            "almost no text could be extracted; if this is a scanned PDF, export a text-based PDF or DOCX")
    return text


def _pdf_text(data: bytes) -> str:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def _docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [f"- {p.text}" if p.style is not None and p.style.name.startswith("List") else p.text
             for p in document.paragraphs]
    for table in document.tables:  # many resume templates lay out sections in tables
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)
    return "\n".join(lines)


def clean_text(text: str) -> str:
    """Normalise whitespace and bullet glyphs; keep line structure."""
    text = text.replace(" ", " ").replace("​", "")
    text = re.sub(r"^[ \t]*[•●▪◦‣∙·]\s*", "- ", text, flags=re.MULTILINE)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def content_hash(text: str) -> str:
    """Stable hash of normalised text, used as a cache key."""
    return hashlib.sha256(re.sub(r"\s+", " ", text).strip().lower().encode()).hexdigest()


_SYSTEM = """You extract structured data from resumes. Rules:
- Copy experience bullets verbatim; do not shorten, merge or improve them.
- skills: languages, frameworks, libraries, methods. tools: platforms, cloud services, databases, dev tools.
- Use canonical names (e.g. "PostgreSQL", "Kubernetes"). Technical items only, no soft skills.
- total_years_experience: sum of professional role durations (exclude education); today is {today}.
- skill_years: only for skills whose usage spans identifiable roles.
- Never invent anything that is not in the resume."""


async def structure_resume(llm: StructuredLLM, storage: Storage | None, text: str,
                           today: str) -> tuple[ResumeStructured, bool]:
    """Return the structured resume and whether it came from cache."""
    key = content_hash(text)
    if storage and (hit := storage.get_cached_parse("resume", key, llm.model)):
        return ResumeStructured.model_validate(hit), True
    resume = await llm.generate_json(
        _SYSTEM.format(today=today), f"RESUME:\n{text}", ResumeStructured, max_tokens=3072)
    if storage:
        storage.put_cached_parse("resume", key, llm.model, resume.model_dump())
    return resume, False


def bullets_with_ids(resume: ResumeStructured) -> list[Bullet]:
    """Flatten experience bullets, most recent role first, ids B1..Bn."""
    out: list[Bullet] = []
    for role in resume.experience:
        label = f"{role.title} @ {role.company}".strip(" @")
        for text in role.bullets:
            if text.strip():
                out.append(Bullet(f"B{len(out) + 1}", label, text.strip()))
    return out
